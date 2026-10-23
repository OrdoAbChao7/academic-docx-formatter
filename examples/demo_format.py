# -*- coding: utf-8 -*-
"""
demo_format.py
Quick start demonstration of academic-docx-formatter.
Tests both presets: 'compact' (10.5pt 五号) and 'comprehensive' (12.0pt 小四).
"""

import os
import sys
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'scripts')))
from format_docx import format_document
from audit_docx import audit_docx

def create_sample_unformatted_docx(filename="sample_unformatted.docx"):
    doc = docx.Document()
    
    # Title Banner
    doc.add_paragraph("2026年第十四届数媒大赛 · 超级个体（OPC）专项赛申报材料")
    # Title
    doc.add_paragraph("微光译语——可穿戴无声交互系统技术说明书")
    # Subtitle
    doc.add_paragraph("（参赛团队匿名盲审申报材料 · 正式版）")
    
    # Metadata callout
    doc.add_paragraph("【项目基本信息】 参赛赛项：超级个体（OPC）· 智能终端与可穿戴消费 | 团队代号：微光译语 | 核心技术：940nm微距光电感知 / ESP32-C3端侧TinyML。")
    
    # Heading 1
    doc.add_paragraph("一、 项目研发背景与核心创新点")
    
    # Body Paragraph
    p_body1 = doc.add_paragraph(
        "言语交流是维系自主生活能力与基本日常意愿表达的关键途径。流行病学统计显示，我国脑卒中患者已超 1300 万人，"
        "其中约 30%~40% 伴随不同程度的失语或构音障碍；ALS 渐冻症患者因延髓运动神经元病变无法正常发声。"
        "传统方案面临严重隐私顾虑与强光干扰痛点，患者迫切需要一种非侵入、低延迟且全脱机的辅助表达系统。"
    )
    p_body1.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY # Causes stretching bug
    
    # Heading 2
    doc.add_paragraph("1.1 关键技术指标与算力权衡")
    
    # Table Caption
    doc.add_paragraph("表 1：系统核心物理指标与实测参数对照表")
    
    # Table
    table = doc.add_table(rows=3, cols=3)
    data = [
        ["指标项名称", "实测参数值", "国家标准依据"],
        ["推理延迟 (Latency)", "1.0 ms", "毫秒级实时响应规范"],
        ["SRAM 占用 (RAM)", "6.7 KB", "微控制器轻量部署标准"]
    ]
    for r_i, row in enumerate(data):
        for c_i, val in enumerate(row):
            table.rows[r_i].cells[c_i].paragraphs[0].text = val

    # Figure Caption
    doc.add_paragraph("图 1：自研 1D-CNN 模型在 300 例微动样本上的混淆矩阵图")
            
    doc.save(filename)
    print(f"Created sample unformatted document: {filename}")
    return filename

if __name__ == "__main__":
    sample_file = "sample_unformatted.docx"
    compact_out = "sample_compact_formatted.docx"
    comp_out = "sample_comprehensive_formatted.docx"
    
    create_sample_unformatted_docx(sample_file)
    
    print("\n>>> AUDITING UNFORMATTED DOCUMENT:")
    audit_docx(sample_file, verbose=True)
    
    print("\n>>> [1/2] FORMATTING PRESET: COMPACT (10.5pt 五号, 竞赛限页申报版)")
    format_document(sample_file, output_path=compact_out, preset="compact")
    audit_docx(compact_out, verbose=False)
    
    print("\n>>> [2/2] FORMATTING PRESET: COMPREHENSIVE (12.0pt 小四, 论文工程详本版)")
    format_document(sample_file, output_path=comp_out, preset="comprehensive")
    audit_docx(comp_out, verbose=False)

    # Clean up
    for f in [sample_file, compact_out, comp_out]:
        if os.path.exists(f): os.remove(f)
    print("\nAll demo runs completed and cleaned up successfully!")
